"""
Single source of truth for all DAX measures (Project 1A).
Generates:  securitisation_measures.dax.md | DAXStudio_Benchmark.dax | 483553A_ManyaSyal_DAX_Measure_Dictionary.xlsx
Run after securitisation_pipeline.py (needs powerbi_data/ExpectedMeasureValues.csv).
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(".")
M = []   # (id, day, category, name, dax, purpose, fmt)
_n = {}


def m(day, cat, name, dax, purpose, fmt=""):
    _n[day] = _n.get(day, 0) + 1
    M.append((f"D{day:02d}-{_n[day]:02d}", day, cat, name, dax.strip("\n"), purpose, fmt))


T, P, Y, S = "FactLoanTape", "FactMonthlyPerformance", "FactDynamicLoss", "FactStaticPool"

# =========================================================================== PARAMETERS (calculated tables)
PARAM_TABLES = [
    ("ParamDefaultMult", "GENERATESERIES ( 0.5, 5, 0.1 )", "Default-rate multiplier (x)"),
    ("ParamRecoveryHaircut", "GENERATESERIES ( 0, 0.5, 0.05 )", "Recovery haircut = additive LGD (pts)"),
    ("ParamCPRChange", "GENERATESERIES ( -0.1, 0.1, 0.01 )", "Change in annual CPR (pts)"),
    ("ParamRateShock", "GENERATESERIES ( 0, 500, 25 )", "Interest-rate shock (bps)"),
    ("ParamStage2DPD", "GENERATESERIES ( 0, 120, 5 )", "Stage 2 DPD threshold (tape = 30, brief = 60)"),
    ("ParamStage3DPD", "GENERATESERIES ( 60, 180, 5 )", "Stage 3 DPD threshold (tape = 90, brief = 120)"),
    ("ParamDelinqTrigger", "GENERATESERIES ( 0.05, 0.25, 0.005 )", "Delinquency trigger level"),
    ("ParamCNLTrigger", "GENERATESERIES ( 0.005, 0.1, 0.0025 )", "Cumulative net loss trigger level"),
    ("ParamOCTrigger", "GENERATESERIES ( 1, 1.3, 0.01 )", "OC ratio trigger (minimum)"),
]

# =========================================================================== DAY 3 - CORE
m(3, "Pool", "Total Outstanding Balance", f"SUM ( {T}[CurrentBalance] )", "Current principal outstanding on the pool.", "#,##0")
m(3, "Pool", "Original Balance", "SUM ( DimLoan[OriginalLoanAmount] )", "Original principal advanced at origination.", "#,##0")
m(3, "Pool", "Pool Factor", "DIVIDE ( [Total Outstanding Balance], [Original Balance] )", "Share of original balance still outstanding.", "0.0%")
m(3, "Pool", "Total EAD", f"SUM ( {T}[EAD] )", "Exposure at default used for ECL.", "#,##0")
m(3, "Counts", "Loan Count", f"COUNTROWS ( {T} )", "All loans on the tape.", "#,##0")
m(3, "Counts", "Active Loans", f"CALCULATE ( COUNTROWS ( {T} ), {T}[CurrentBalance] > 0 )", "Loans with a positive balance (24 loans carry zero balance).", "#,##0")
m(3, "Counts", "Defaulted Loans", f"CALCULATE ( COUNTROWS ( {T} ), {T}[IsDefaulted] = 1 )", "Loans flagged Default or Repossessed.", "#,##0")
m(3, "Counts", "Prepaid Loans", f"CALCULATE ( COUNTROWS ( {T} ), {T}[PrepaymentAmount] > 0 )", "Loans with a voluntary prepayment.", "#,##0")
m(3, "Balances", "Defaulted Balance", f"CALCULATE ( [Total Outstanding Balance], {T}[IsDefaulted] = 1 )", "Outstanding balance of defaulted loans.", "#,##0")
m(3, "Delinquency", "Balance Current", f'CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyStatus] = "Current" )', "Balance with no arrears.", "#,##0")
m(3, "Delinquency", "Balance 1-29 DPD", f'CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyStatus] = "1-29 DPD" )', "Balance 1-29 days past due.", "#,##0")
m(3, "Delinquency", "Balance 30 DPD", f'CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyStatus] = "30-59 DPD" )', "Balance 30-59 days past due.", "#,##0")
m(3, "Delinquency", "Balance 60 DPD", f'CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyStatus] = "60-89 DPD" )', "Balance 60-89 days past due.", "#,##0")
m(3, "Delinquency", "Balance 90 DPD", f'CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyStatus] = "90-119 DPD" )', "Balance 90-119 days past due.", "#,##0")
m(3, "Delinquency", "Balance 120+ DPD", f'CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyStatus] IN {{ "120+ DPD", "Default", "Repossessed" }} )', "Balance 120+ DPD plus Default/Repossessed.", "#,##0")
m(3, "Delinquency", "Delinquent 30+ Balance", f"CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyDays] >= 30 )", "Balance 30 or more days past due.", "#,##0")
m(3, "Delinquency", "Delinquent 30+ %", "DIVIDE ( [Delinquent 30+ Balance], [Total Outstanding Balance] )", "30+ DPD as % of pool balance (SMA-1 and worse).", "0.0%")
m(3, "Delinquency", "Delinquent 60+ Balance", f"CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyDays] >= 60 )", "Balance 60 or more days past due.", "#,##0")
m(3, "Delinquency", "Delinquent 60+ %", "DIVIDE ( [Delinquent 60+ Balance], [Total Outstanding Balance] )", "60+ DPD as % of pool; drives the delinquency trigger.", "0.0%")
m(3, "Delinquency", "Delinquent 90+ Balance", f"CALCULATE ( [Total Outstanding Balance], {T}[DelinquencyDays] >= 90 )", "Balance 90 or more days past due (NPA).", "#,##0")
m(3, "Delinquency", "Delinquent 90+ %", "DIVIDE ( [Delinquent 90+ Balance], [Total Outstanding Balance] )", "NPA ratio.", "0.0%")
_w = lambda col, tbl=T: f"DIVIDE ( SUMX ( {T}, {tbl}[{col}] * {T}[CurrentBalance] ), [Total Outstanding Balance] )"
m(3, "Weighted avg", "WAC", f"DIVIDE ( SUMX ( {T}, RELATED ( DimLoan[InterestRate] ) * {T}[CurrentBalance] ), [Total Outstanding Balance] ) / 100", "Weighted-average coupon (rate in tape is in % points).", "0.00%")
m(3, "Weighted avg", "WAM", _w("RemainingTerm"), "Weighted-average remaining maturity (months).", "0.0")
m(3, "Weighted avg", "WALA", _w("MonthsOnBook"), "Weighted-average loan age (months on book).", "0.0")
m(3, "Weighted avg", "WALT", f"DIVIDE ( SUMX ( {T}, RELATED ( DimLoan[OriginalTerm] ) * {T}[CurrentBalance] ), [Total Outstanding Balance] )", "Weighted-average original loan term (months).", "0.0")
m(3, "Weighted avg", "WA LTV Origination", _w("LTV_AtOrigination"), "Balance-weighted LTV at origination.", "0.0%")
m(3, "Weighted avg", "WA LTV Current", _w("LTV_Current"), "Balance-weighted current LTV.", "0.0%")
m(3, "Weighted avg", "WA CIBIL", _w("CIBIL_Score_Current"), "Balance-weighted current CIBIL score.", "0")
m(3, "Weighted avg", "WA CIBIL Origination", _w("CIBIL_Score_Origination"), "Balance-weighted CIBIL at origination.", "0")
m(3, "Weighted avg", "WA DTI", _w("DTI_Ratio"), "Balance-weighted debt-to-income.", "0.0%")
m(3, "Pool", "Average Loan Balance", "DIVIDE ( [Total Outstanding Balance], [Active Loans] )", "Mean balance of active loans.", "#,##0")
m(3, "Pool", "Largest Loan Balance", f"MAX ( {T}[CurrentBalance] )", "Single largest exposure.", "#,##0")
m(3, "Pool", "Insured Loan %", "DIVIDE ( CALCULATE ( [Total Outstanding Balance], DimLoan[HasInsurance] = 1 ), [Total Outstanding Balance] )", "Balance share with any vehicle insurance.", "0.0%")
m(3, "Pool", "New Vehicle %", "DIVIDE ( CALCULATE ( [Total Outstanding Balance], DimVehicle[IsNewVehicle] = 1 ), [Total Outstanding Balance] )", "Balance share financing new vehicles.", "0.0%")
m(3, "Pool", "Modified Loan %", f"DIVIDE ( CALCULATE ( COUNTROWS ( {T} ), DimLoan[IsModified] = 1 ), [Loan Count] )", "Share of loans restructured (deferral, extension, rate cut).", "0.0%")
m(3, "Loss", "Total Prepayments", f"SUM ( {T}[PrepaymentAmount] )", "Voluntary prepayments recorded on the tape.", "#,##0")
m(3, "Loss", "Total Loss Amount", f"SUM ( {T}[LossAmount] )", "Gross realised loss on the tape.", "#,##0")
m(3, "Loss", "Total Recoveries", f"SUM ( {T}[RecoveryAmount] )", "Recoveries on the tape.", "#,##0")
m(3, "Loss", "Total Net Loss (Tape)", f"SUM ( {T}[NetLoss] )", "Net realised loss on the tape.", "#,##0")
m(3, "Collections", "Payment Collection Rate", f"DIVIDE ( SUM ( {T}[LastPaymentAmount] ), SUM ( {T}[ScheduledPaymentDue] ) )", "Last payment received vs scheduled payment due.", "0.0%")

# =========================================================================== DAY 4 - TIME INTELLIGENCE (FactDynamicLoss <-> DimDate on MonthEnd)
yc = lambda col: f"CALCULATE ( SUM ( {Y}[{col}] ), {Y}[IncludeInTrend] = 1 )"
m(4, "Base", "Net Loss (Monthly)", yc("NetLoss_ThisMonth"), "Monthly net loss (second 30-May row excluded - see Data Quality DQ08).", "#,##0")
m(4, "Base", "Gross Loss (Monthly)", yc("GrossLoss_ThisMonth"), "Monthly gross loss.", "#,##0")
m(4, "Base", "Recoveries (Monthly)", yc("Recoveries_ThisMonth"), "Monthly recoveries.", "#,##0")
m(4, "Base", "New Defaults Balance (Monthly)", yc("NewDefaults_Balance"), "Balance newly defaulted in the month.", "#,##0")
m(4, "Base", "BOP Balance (Monthly)", yc("BOP_Balance"), "Beginning-of-period pool balance.", "#,##0")
m(4, "Base", "Monthly Default Rate", "DIVIDE ( [New Defaults Balance (Monthly)], [BOP Balance (Monthly)] )", "New defaults / opening balance.", "0.000%")
m(4, "Base", "Monthly Net Loss Rate", "DIVIDE ( [Net Loss (Monthly)], [BOP Balance (Monthly)] )", "Net loss / opening balance.", "0.000%")
m(4, "To-date", "Net Loss MTD", "TOTALMTD ( [Net Loss (Monthly)], DimDate[Date] )", "Month-to-date net loss.", "#,##0")
m(4, "To-date", "Net Loss QTD", "TOTALQTD ( [Net Loss (Monthly)], DimDate[Date] )", "Quarter-to-date net loss.", "#,##0")
m(4, "To-date", "Net Loss YTD", "TOTALYTD ( [Net Loss (Monthly)], DimDate[Date] )", "Calendar year-to-date net loss.", "#,##0")
m(4, "To-date", "Net Loss FYTD", """
VAR fy = MAX ( DimDate[FiscalYear] )
VAR d  = MAX ( DimDate[Date] )
RETURN CALCULATE ( [Net Loss (Monthly)], FILTER ( ALL ( DimDate ), DimDate[FiscalYear] = fy && DimDate[Date] <= d ) )""", "Indian fiscal year-to-date (April-March) net loss; built on DimDate[FiscalYear] so it does not depend on locale date strings.", "#,##0")
m(4, "Comparison", "Net Loss Prior Month", "CALCULATE ( [Net Loss (Monthly)], DATEADD ( DimDate[Date], -1, MONTH ) )", "Net loss one month earlier.", "#,##0")
m(4, "Comparison", "Net Loss MoM %", "DIVIDE ( [Net Loss (Monthly)] - [Net Loss Prior Month], [Net Loss Prior Month] )", "Month-on-month change in net loss.", "0.0%")
m(4, "Comparison", "Default Rate Prior Quarter", "CALCULATE ( [Monthly Default Rate], DATEADD ( DimDate[Date], -1, QUARTER ) )", "Monthly default rate one quarter earlier.", "0.000%")
m(4, "Comparison", "Default Rate QoQ Change", "[Monthly Default Rate] - [Default Rate Prior Quarter]", "Quarter-on-quarter change (pts).", "0.000%")
m(4, "Comparison", "Default Rate vs Prior Year", "CALCULATE ( [Monthly Default Rate], SAMEPERIODLASTYEAR ( DimDate[Date] ) )", "Prior-year comparator (blank: file covers < 12 months).", "0.000%")
_last = f"MAX ( {Y}[MonthEnd] )"
m(4, "Rolling", "Rolling 3M Default Rate", f'''
VAR LastD = {_last}
VAR Win = DATESINPERIOD ( DimDate[Date], LastD, -3, MONTH )
RETURN DIVIDE ( CALCULATE ( [New Defaults Balance (Monthly)], Win ), CALCULATE ( [BOP Balance (Monthly)], Win ) )''', "Balance-weighted default rate over the trailing 3 months.", "0.000%")
m(4, "Rolling", "Rolling 12M Default Rate", f'''
VAR LastD = {_last}
VAR Win = DATESINPERIOD ( DimDate[Date], LastD, -12, MONTH )
RETURN DIVIDE ( CALCULATE ( [New Defaults Balance (Monthly)], Win ), CALCULATE ( [BOP Balance (Monthly)], Win ) )''', "Balance-weighted default rate over the trailing 12 months.", "0.000%")
m(4, "Rolling", "Rolling 3M Net Loss Rate", f'''
VAR LastD = {_last}
VAR Win = DATESINPERIOD ( DimDate[Date], LastD, -3, MONTH )
RETURN DIVIDE ( CALCULATE ( [Net Loss (Monthly)], Win ), CALCULATE ( [BOP Balance (Monthly)], Win ) )''', "Trailing 3-month net loss rate.", "0.000%")
_months12 = f'''
VAR LastD = {_last}
VAR Months12 =
    CALCULATETABLE ( VALUES ( DimDate[Date] ), DATESINPERIOD ( DimDate[Date], LastD, -12, MONTH ), DimDate[IsMonthEnd] = 1 )'''
m(4, "Alerts", "Default Rate 12M Average", _months12 + "\nRETURN AVERAGEX ( Months12, [Monthly Default Rate] )", "Mean monthly default rate over 12 months.", "0.000%")
m(4, "Alerts", "Default Rate 12M StdDev", _months12 + "\nRETURN STDEVX.P ( Months12, [Monthly Default Rate] )", "Standard deviation of monthly default rate.", "0.000%")
m(4, "Alerts", "Default Rate Z-Score", "DIVIDE ( [Monthly Default Rate] - [Default Rate 12M Average], [Default Rate 12M StdDev] )", "Standard deviations from the 12-month average.", "0.00")
m(4, "Alerts", "Default Rate Trend Alert", 'IF ( ABS ( [Default Rate Z-Score] ) > 2, "ALERT: > 2 std dev", "Within range" )', "Flags metrics moving more than 2 std dev from their 12M average (brief Day 4).", "")
m(4, "Alerts", "Delinquency 30+ Z-Score", f'''
VAR LastD = {_last}
VAR Months12 =
    CALCULATETABLE ( VALUES ( DimDate[Date] ), DATESINPERIOD ( DimDate[Date], LastD, -12, MONTH ), DimDate[IsMonthEnd] = 1 )
VAR Cur = [DPD 30+ % (Date View)]
RETURN DIVIDE ( Cur - AVERAGEX ( Months12, [DPD 30+ % (Date View)] ), STDEVX.P ( Months12, [DPD 30+ % (Date View)] ) )''', "Z-score of 30+ DPD share vs its 12M history.", "0.00")

# =========================================================================== DAY 5 - ADVANCED / VIRTUAL TABLES
m(5, "Strat", "Pool Share %", "DIVIDE ( [Total Outstanding Balance], CALCULATE ( [Total Outstanding Balance], ALLSELECTED () ) )", "Share of the selected pool - use with any stratification axis (LTV band, CIBIL band, vintage).", "0.0%")
m(5, "Strat", "LTV > 90% Share", f"DIVIDE ( CALCULATE ( [Total Outstanding Balance], {T}[LTV_Current] > 0.9 ), [Total Outstanding Balance] )", "Balance share with current LTV above 90%.", "0.0%")
m(5, "Strat", "CIBIL < 650 Share", f"DIVIDE ( CALCULATE ( [Total Outstanding Balance], {T}[CIBIL_Score_Current] < 650 ), [Total Outstanding Balance] )", "Balance share with sub-650 CIBIL.", "0.0%")
m(5, "Concentration", "Top 3 States", '''
CONCATENATEX (
    TOPN ( 3, VALUES ( DimGeography[State] ), [Total Outstanding Balance], DESC ),
    DimGeography[State] & " " & FORMAT ( DIVIDE ( [Total Outstanding Balance], CALCULATE ( [Total Outstanding Balance], ALLSELECTED ( DimGeography ) ) ), "0.0%" ),
    ", ", [Total Outstanding Balance], DESC )''', "Text label of the three largest states with their share.", "")
m(5, "Concentration", "Top 5 State Concentration", "DIVIDE ( SUMX ( TOPN ( 5, VALUES ( DimGeography[State] ), [Total Outstanding Balance], DESC ), [Total Outstanding Balance] ), [Total Outstanding Balance] )", "Share of the pool in the five largest states (TOPN virtual table).", "0.0%")
m(5, "Concentration", "Top 10 Loan Concentration", "DIVIDE ( SUMX ( TOPN ( 10, VALUES ( DimLoan[LoanID] ), [Total Outstanding Balance], DESC ), [Total Outstanding Balance] ), [Total Outstanding Balance] )", "Share held by the ten largest loans.", "0.0%")
m(5, "Concentration", "Top 10 Borrower Concentration", "DIVIDE ( SUMX ( TOPN ( 10, VALUES ( DimBorrower[BorrowerID] ), [Total Outstanding Balance], DESC ), [Total Outstanding Balance] ), [Total Outstanding Balance] )", "Share held by the ten largest borrowers.", "0.0%")
_hhi = lambda col: f"VAR Tot = [Total Outstanding Balance]\nRETURN SUMX ( VALUES ( {col} ), POWER ( DIVIDE ( [Total Outstanding Balance], Tot ), 2 ) )"
m(5, "Concentration", "HHI State", _hhi("DimGeography[State]"), "Herfindahl-Hirschman Index by state (0-1).", "0.0000")
m(5, "Concentration", "HHI Servicer", _hhi("DimLoan[ServicerName]"), "HHI by servicer.", "0.0000")
m(5, "Concentration", "HHI Vehicle Make", _hhi("DimVehicle[VehicleMake]"), "HHI by vehicle manufacturer.", "0.0000")
m(5, "Concentration", "HHI Borrower", _hhi("DimBorrower[BorrowerID]"), "HHI by borrower (borrower IDs are unique in this tape).", "0.0000")
m(5, "Ranking", "State Rank", "RANKX ( ALLSELECTED ( DimGeography[State] ), [Total Outstanding Balance], , DESC, DENSE )", "Rank of each state by balance.", "0")
m(5, "Ranking", "Distinct States", "DISTINCTCOUNT ( DimLoan[State] )", "Number of states in the selection.", "0")
m(5, "Context transition", "Portfolio ECL (Context Transition)", f'''
SUMX (
    {T},
    CALCULATE ( AVERAGE ( {T}[PD_Estimate] ) * AVERAGE ( {T}[LGD_Estimate] ) ) * {T}[EAD]
)''', "Context-transition demonstration: CALCULATE inside SUMX turns each row into a filter context (brief Day 5).", "#,##0")

# =========================================================================== DAY 6 - WATERFALL & TRIGGERS
pp = lambda col: f"VAR p = MAX ( FactPool[period] )\nRETURN CALCULATE ( SUM ( FactPool[{col}] ), FactPool[period] = p )"
m(6, "Pool cash flow", "Ending Pool Balance", pp("ending_pool"), "Pool balance at the last period in context.", "#,##0")
m(6, "Pool cash flow", "Interest Collected", "SUM ( FactPool[interest_collected] )", "Interest collections.", "#,##0")
m(6, "Pool cash flow", "Servicing Fees", "SUM ( FactPool[servicing_fees] )", "Senior servicing fee (0.5% p.a., assumption).", "#,##0")
m(6, "Pool cash flow", "Net Loss (Projected)", "SUM ( FactPool[net_loss] )", "Projected net loss under the scenario.", "#,##0")
m(6, "Pool cash flow", "Excess Spread", "SUM ( FactPool[excess_spread] )", "Interest left after fees and note coupons.", "#,##0")
m(6, "Pool cash flow", "Excess Spread % (annualised)", "DIVIDE ( SUM ( FactPool[excess_spread] ), SUM ( FactPool[beginning_pool] ) ) * 12", "Annualised excess spread / pool balance.", "0.00%")
m(6, "Pool cash flow", "Cumulative Loss % (Projected)", "VAR p = MAX ( FactPool[period] )\nRETURN CALCULATE ( MAX ( FactPool[cumulative_loss_pct] ), FactPool[period] = p )", "Cumulative projected net loss / opening pool at the period in context.", "0.00%")
m(6, "Pool cash flow", "OC Ratio", "VAR p = MAX ( FactPool[period] )\nRETURN CALCULATE ( MAX ( FactPool[oc_ratio] ), FactPool[period] = p )", "Pool balance / (Class A+B+C notes).", "0.000")
m(6, "Tranche", "Tranche Beginning Balance", "VAR p = MIN ( FactWaterfall[period] )\nRETURN CALCULATE ( SUM ( FactWaterfall[beginning_balance] ), FactWaterfall[period] = p )", "Opening note balance for the first period in context.", "#,##0")
m(6, "Tranche", "Tranche Ending Balance", "VAR p = MAX ( FactWaterfall[period] )\nRETURN CALCULATE ( SUM ( FactWaterfall[ending_balance] ), FactWaterfall[period] = p )", "Closing note balance at the last period in context.", "#,##0")
m(6, "Tranche", "Interest Due", "SUM ( FactWaterfall[interest_due] )", "Coupon due.", "#,##0")
m(6, "Tranche", "Interest Paid", "SUM ( FactWaterfall[interest_paid] )", "Coupon paid out of available interest.", "#,##0")
m(6, "Tranche", "Interest Shortfall", "[Interest Due] - [Interest Paid]", "Unpaid coupon.", "#,##0")
m(6, "Tranche", "Principal Paid", "SUM ( FactWaterfall[principal_paid] )", "Principal distributed (sequential A > B > C > Equity).", "#,##0")
m(6, "Tranche", "Tranche Writedown", "SUM ( FactWaterfall[writedown] )", "Losses allocated to the tranche (Equity first, Class A last).", "#,##0")
m(6, "Tranche", "Writedown % of Tranche", "DIVIDE ( [Tranche Writedown], CALCULATE ( SUM ( FactWaterfall[beginning_balance] ), FactWaterfall[period] = 1 ) )", "Share of original tranche balance written down.", "0.0%")
m(6, "Tranche", "Credit Enhancement %", "MAX ( DimTranche[attachment] )", "Subordination below the tranche (attachment point).", "0.0%")
m(6, "Tranche", "Tranche Loss % (Attach-Detach)", '''
VAR CumLoss = [Cumulative Loss % (Projected)]
VAR Attach  = MAX ( DimTranche[attachment] )
VAR Detach  = MAX ( DimTranche[detachment] )
RETURN DIVIDE ( MAX ( 0, MIN ( CumLoss, Detach ) - Attach ), Detach - Attach )''', "Closed-form tranche loss from pool loss and attachment/detachment points; cross-checks the cash-flow write-down.", "0.0%")
m(6, "Tranche", "Principal Available", "SUM ( FactPool[scheduled_principal] ) + SUM ( FactPool[prepayments] ) + SUM ( FactPool[recoveries] ) + SUM ( FactPool[excess_spread_used_for_losses] )", "Cash available for sequential principal pay-down.", "#,##0")
m(6, "Tranche", "Principal to Class A", '''
VAR BalA = CALCULATE ( SUM ( FactWaterfall[beginning_balance] ), DimTranche[tranche] = "Class A (Senior)" )
RETURN MIN ( [Principal Available], BalA )''', "Sequential allocation: senior class receives principal first.", "#,##0")
m(6, "Tranche", "Principal to Class B", '''
VAR BalB = CALCULATE ( SUM ( FactWaterfall[beginning_balance] ), DimTranche[tranche] = "Class B (Mezzanine)" )
RETURN MIN ( MAX ( [Principal Available] - [Principal to Class A], 0 ), BalB )''', "Mezzanine receives what remains after Class A.", "#,##0")
m(6, "Tranche", "Tranche Status", '''
SWITCH ( TRUE (),
    [Writedown % of Tranche] = 0, "Intact",
    [Writedown % of Tranche] < 0.5, "Impaired",
    "Wiped out" )''', "Intact / impaired / wiped out label.", "")
for nm, tbl in [("Delinq Trigger Value", ("ParamDelinqTrigger", 0.12)), ("CNL Trigger Value", ("ParamCNLTrigger", 0.02)), ("OC Trigger Value", ("ParamOCTrigger", 1.05))]:
    m(6, "Triggers", nm, f"SELECTEDVALUE ( {tbl[0]}[Value], {tbl[1]} )", f"What-if driven {nm.replace(' Value', '').lower()} level (default is an ASSUMPTION).", "0.0%" if "OC" not in nm else "0.00")
m(6, "Triggers", "Delinquency Trigger Status", 'IF ( [Delinquent 60+ %] > [Delinq Trigger Value], "BREACH", "OK" )', "Tape 60+ DPD % vs delinquency trigger.", "")
m(6, "Triggers", "Delinquency Trigger Headroom", "[Delinq Trigger Value] - [Delinquent 60+ %]", "Trigger level minus actual (negative = breach).", "0.0%")
m(6, "Triggers", "CNL Trigger Status", 'IF ( [Cumulative Loss % (Projected)] > [CNL Trigger Value], "BREACH", "OK" )', "Projected cumulative net loss vs trigger.", "")
m(6, "Triggers", "OC Trigger Status", 'IF ( [OC Ratio] < [OC Trigger Value], "BREACH", "OK" )', "OC ratio vs minimum.", "")
m(6, "Triggers", "Trigger Breach Count", '( [Delinquency Trigger Status] = "BREACH" ) + ( [CNL Trigger Status] = "BREACH" ) + ( [OC Trigger Status] = "BREACH" )', "Number of triggers breached.", "0")
m(6, "Triggers", "CNL Trigger Breach Month", '''
VAR t = [CNL Trigger Value]
RETURN CALCULATE ( MIN ( DimPeriod[period] ), FILTER ( ALL ( DimPeriod[period] ), [Cumulative Loss % (Projected)] > t ) )''', "First projected month in which cumulative loss exceeds the trigger (blank = never).", "0")

# =========================================================================== DAY 7 - DPD & ROLL RATES
nopre = f"{P}[PreOrigination] = 0"
m(7, "Distribution", "Snapshot Loans", f"CALCULATE ( DISTINCTCOUNT ( {P}[LoanID] ), {nopre} )", "Loans with a post-origination snapshot (use SnapshotLabel on the axis).", "#,##0")
m(7, "Distribution", "Snapshot Balance", f"CALCULATE ( SUM ( {P}[CurrentBalance] ), {nopre} )", "Balance on the snapshots in context (pre-origination rows excluded).", "#,##0")
m(7, "Distribution", "DPD 30+ Balance (Snapshot)", f"CALCULATE ( [Snapshot Balance], {P}[BucketOrder] >= 2 )", "30+ DPD balance by snapshot.", "#,##0")
m(7, "Distribution", "DPD 30+ %", "DIVIDE ( [DPD 30+ Balance (Snapshot)], [Snapshot Balance] )", "30+ DPD share by snapshot (axis = SnapshotLabel).", "0.0%")
m(7, "Distribution", "DPD 60+ %", f"DIVIDE ( CALCULATE ( [Snapshot Balance], {P}[BucketOrder] >= 3 ), [Snapshot Balance] )", "60+ DPD share by snapshot.", "0.0%")
m(7, "Distribution", "DPD 90+ %", f"DIVIDE ( CALCULATE ( [Snapshot Balance], {P}[BucketOrder] >= 4 ), [Snapshot Balance] )", "90+ DPD share by snapshot.", "0.0%")
m(7, "Distribution", "Bucket Balance %", "DIVIDE ( [Snapshot Balance], CALCULATE ( [Snapshot Balance], ALL ( DimBucket ) ) )", "Share of snapshot balance by DPD bucket.", "0.0%")
m(7, "Distribution", "DPD 30+ % (Date View)", f"CALCULATE ( [DPD 30+ %], {P}[IncludeInTrend] = 1 )", "30+ DPD share on a calendar axis (duplicate May snapshot excluded).", "0.0%")
m(7, "Roll rates", "Transition Count", f"CALCULATE ( COUNTROWS ( {P} ), {P}[ValidTransition] = 1 )", "Loan-month transitions between consecutive snapshots.", "#,##0")
m(7, "Roll rates", "Roll Rate %", f"DIVIDE ( [Transition Count], CALCULATE ( [Transition Count], ALL ( DimBucket[DPD_Bucket] ) ) )", "Matrix measure: rows = {P}[PriorBucketSeq], columns = DimBucket[DPD_Bucket]; % of the prior bucket moving to each bucket.".replace("{P}", P), "0.0%")
m(7, "Roll rates", "Worsened Transitions", f"CALCULATE ( [Transition Count], FILTER ( {P}, {P}[BucketOrder] > {P}[PriorBucketOrder] ) )", "Transitions to a worse bucket.", "#,##0")
m(7, "Roll rates", "Cured Transitions", f"CALCULATE ( [Transition Count], FILTER ( {P}, {P}[BucketOrder] < {P}[PriorBucketOrder] ) )", "Transitions to a better bucket.", "#,##0")
m(7, "Roll rates", "Roll Forward Rate", "DIVIDE ( [Worsened Transitions], [Transition Count] )", "Share of transitions that worsen.", "0.0%")
m(7, "Roll rates", "Cure Rate", "DIVIDE ( [Cured Transitions], [Transition Count] )", "Share of transitions that improve.", "0.0%")
m(7, "Roll rates", "Cure Rate (30+ DPD)", f"DIVIDE ( CALCULATE ( [Cured Transitions], {P}[PriorBucketOrder] >= 2 ), CALCULATE ( [Transition Count], {P}[PriorBucketOrder] >= 2 ) )", "Cure rate for loans already 30+ DPD (zero in this data).", "0.0%")
_rr = lambda fr, to: f'DIVIDE ( CALCULATE ( [Transition Count], {P}[PriorBucketSeq] = "{fr}", {to} ), CALCULATE ( [Transition Count], {P}[PriorBucketSeq] = "{fr}" ) )'
m(7, "Roll rates", "Roll Rate Current to Current", _rr("Current", f'{P}[DPD_Bucket] = "Current"'), "Current loans staying current.", "0.0%")
m(7, "Roll rates", "Roll Rate Current to 1-29", _rr("Current", f'{P}[DPD_Bucket] = "1-29 DPD"'), "Current loans slipping into 1-29 DPD.", "0.0%")
m(7, "Roll rates", "Roll Rate 30-59 to 60-89", _rr("30-59 DPD", f'{P}[DPD_Bucket] = "60-89 DPD"'), "30-59 DPD loans rolling to 60-89.", "0.0%")
m(7, "Roll rates", "Roll Rate 60-89 to 90+", _rr("60-89 DPD", f"{P}[BucketOrder] >= 4"), "60-89 DPD loans rolling into 90+ / default.", "0.0%")
_age = lambda lo, hi: f"CALCULATE ( [Snapshot Balance], {P}[ConsecutiveMonthsDelinquent] >= {lo}" + (f", {P}[ConsecutiveMonthsDelinquent] <= {hi}" if hi else "") + " )"
m(7, "Ageing", "Balance Delinquent 1-2 Months", _age(1, 2), "Balance delinquent 1-2 consecutive months.", "#,##0")
m(7, "Ageing", "Balance Delinquent 3-5 Months", _age(3, 5), "Balance delinquent 3-5 consecutive months.", "#,##0")
m(7, "Ageing", "Balance Delinquent 6+ Months", _age(6, None), "Balance delinquent 6+ consecutive months.", "#,##0")
m(7, "Ageing", "Avg Months Delinquent", f"CALCULATE ( AVERAGE ( {P}[ConsecutiveMonthsDelinquent] ), {P}[ConsecutiveMonthsDelinquent] > 0, {nopre} )", "Average consecutive months delinquent among delinquent loans.", "0.0")
m(7, "Ageing", "Loans Delinquent 3+ Months (Latest)", f"VAR s = MAX ( {P}[SnapshotSeq] )\nRETURN CALCULATE ( DISTINCTCOUNT ( {P}[LoanID] ), {P}[SnapshotSeq] = s, {P}[ConsecutiveMonthsDelinquent] >= 3 )", "Loans in arrears 3+ months at the latest snapshot.", "#,##0")

# =========================================================================== DAY 8 - VINTAGE
m(8, "Vintage", "Vintage Original Balance", f"CALCULATE ( SUM ( {S}[OriginalPoolBalance] ), {S}[MonthsOnBook] = 0 )", "Original pool balance (counted once per vintage at MOB 0).", "#,##0")
m(8, "Vintage", "CNL Rate @ MOB", f"DIVIDE ( SUM ( {S}[CumulativeNetLoss] ), SUM ( {S}[OriginalPoolBalance] ) )", "Cumulative net loss rate - put MonthsOnBook on the axis (one MOB per point).", "0.00%")
m(8, "Vintage", "Cumulative Default Rate @ MOB", f"DIVIDE ( SUM ( {S}[CumulativeDefaults_Balance] ), SUM ( {S}[OriginalPoolBalance] ) )", "Cumulative default rate by months on book.", "0.00%")
m(8, "Vintage", "Cumulative Gross Loss Rate @ MOB", f"DIVIDE ( SUM ( {S}[CumulativeGrossLoss] ), SUM ( {S}[OriginalPoolBalance] ) )", "Cumulative gross loss rate by months on book.", "0.00%")
m(8, "Vintage", "Cumulative Prepayment Rate @ MOB", f"DIVIDE ( SUM ( {S}[CumulativePrepayments] ), SUM ( {S}[OriginalPoolBalance] ) )", "Cumulative prepayment rate by months on book.", "0.00%")
m(8, "Vintage", "Pool Factor @ MOB", f"DIVIDE ( SUM ( {S}[RemainingPoolBalance] ), SUM ( {S}[OriginalPoolBalance] ) )", "Remaining balance / original by months on book.", "0.0%")
m(8, "Vintage", "Latest MOB", f"MAX ( {S}[MonthsOnBook] )", "Most recent months-on-book observed.", "0")
m(8, "Vintage", "CNL Rate (Latest MOB)", f'''
DIVIDE (
    SUMX ( VALUES ( {S}[VintageID] ),
        VAR m = CALCULATE ( MAX ( {S}[MonthsOnBook] ) )
        RETURN CALCULATE ( SUM ( {S}[CumulativeNetLoss] ), {S}[MonthsOnBook] = m ) ),
    [Vintage Original Balance] )''', "Cumulative net loss to date, each vintage at its own latest MOB (virtual-table iteration).", "0.00%")
m(8, "Vintage", "Loss Timing % (of CNL to date)", "DIVIDE ( [CNL Rate @ MOB], [CNL Rate (Latest MOB)] )", "Share of loss-to-date incurred by each MOB - the loss timing curve rating agencies ask for.", "0.0%")
m(8, "Vintage", "Marginal Loss Rate (Wtd)", f"DIVIDE ( SUMX ( {S}, {S}[MarginalLossRate] * {S}[RemainingPoolBalance] ), SUM ( {S}[RemainingPoolBalance] ) )", "Monthly marginal loss rate weighted by remaining balance.", "0.000%")
m(8, "Vintage", "Vintage 30+ Delinquency (Wtd)", f"DIVIDE ( SUMX ( {S}, {S}[CurrentDelinq30Plus] * {S}[RemainingPoolBalance] ), SUM ( {S}[RemainingPoolBalance] ) )", "30+ delinquency weighted by remaining balance.", "0.00%")
_mature = f'FILTER ( ALL ( {S}[VintageID] ), {S}[VintageID] IN {{ "2021-Q1", "2021-Q2", "2021-Q3", "2021-Q4" }} )'
m(8, "Vintage", "Projected Ultimate CNL (MOB 36)", f'''
VAR MaxMOB  = MAX ( {S}[MonthsOnBook] )
VAR Cur     = CALCULATE ( [CNL Rate @ MOB], {S}[MonthsOnBook] = MaxMOB )
VAR RefNow  = CALCULATE ( [CNL Rate @ MOB], {S}[MonthsOnBook] = MaxMOB, {_mature} )
VAR Ref36   = CALCULATE ( [CNL Rate @ MOB], {S}[MonthsOnBook] = 36, {_mature} )
RETURN IF ( MaxMOB >= 36, Cur, DIVIDE ( Cur * Ref36, RefNow ) )''', "Chain-ladder style projection: scales the vintage's CNL by the mature-vintage (2021) development from its MOB to MOB 36.", "0.00%")
m(8, "Vintage", "Vintage WAC", "[WAC]", "Vintage-level WAC: filter flows DimVintage > DimLoan > tape, so the base measure is reused.", "0.00%")
m(8, "Vintage", "Vintage WALA", "[WALA]", "Vintage-level WALA via the same filter path.", "0.0")

# =========================================================================== DAY 9 - IFRS 9 ECL
m(9, "ECL (tape)", "Total ECL (Tape)", f"SUM ( {T}[ECL_Provision] )", "Provision as supplied on the tape.", "#,##0")
m(9, "ECL (tape)", "ECL Coverage %", "DIVIDE ( [Total ECL (Tape)], [Total EAD] )", "Provision / EAD.", "0.00%")
for k in (1, 2, 3):
    m(9, "ECL (tape)", f"ECL Stage {k}", f"CALCULATE ( [Total ECL (Tape)], {T}[IFRS9_Stage] = {k} )", f"Provision on Stage {k} loans.", "#,##0")
for k in (1, 2, 3):
    m(9, "ECL (tape)", f"EAD Stage {k}", f"CALCULATE ( [Total EAD], {T}[IFRS9_Stage] = {k} )", f"EAD of Stage {k} loans.", "#,##0")
m(9, "ECL (tape)", "Stage Mix % (EAD)", "DIVIDE ( [Total EAD], CALCULATE ( [Total EAD], ALL ( DimStage ) ) )", "Stage share of EAD (DimStage on axis).", "0.0%")
m(9, "ECL (tape)", "Avg PD (EAD-weighted)", f"DIVIDE ( SUMX ( {T}, {T}[PD_Estimate] * {T}[EAD] ), [Total EAD] )", "EAD-weighted PD.", "0.00%")
m(9, "ECL (tape)", "Avg LGD (EAD-weighted)", f"DIVIDE ( SUMX ( {T}, {T}[LGD_Estimate] * {T}[EAD] ), [Total EAD] )", "EAD-weighted LGD.", "0.0%")
m(9, "ECL (tape)", "ECL Recomputed (PD x LGD x EAD)", f"SUMX ( {T}, {T}[PD_Estimate] * {T}[LGD_Estimate] * {T}[EAD] )", "Independent loan-level recomputation of the provision.", "#,##0")
m(9, "ECL (tape)", "ECL Recompute Variance", "[ECL Recomputed (PD x LGD x EAD)] - [Total ECL (Tape)]", "Should be zero - tape ECL = PD x LGD x EAD.", "#,##0")
m(9, "LGD", "Realised Recovery Rate (Tape)", f"DIVIDE ( SUM ( {T}[RecoveryAmount] ), SUM ( {T}[LossAmount] ) )", "Recoveries / gross loss on the tape (2 loans only).", "0.0%")
m(9, "LGD", "Net Loss Severity (Dynamic)", f"DIVIDE ( SUM ( {Y}[NetLoss_ThisMonth] ), SUM ( {Y}[NewDefaults_Balance] ) )", "Net loss / defaulted balance over 12 months - the pool LGD used by the challenger model.", "0.0%")
m(9, "LGD", "Gross Loss Severity (Dynamic)", f"DIVIDE ( SUM ( {Y}[GrossLoss_ThisMonth] ), SUM ( {Y}[NewDefaults_Balance] ) )", "Gross loss / defaulted balance.", "0.0%")
m(9, "PD", "PD Term Structure", "MAX ( FactPDTermStructure[CumulativePD] )", "Cumulative PD from the empirical transition matrix (axis: HorizonMonths, legend: DPD_Bucket).", "0.0%")
m(9, "PD", "12M PD from Transition Matrix", "CALCULATE ( MAX ( FactPDTermStructure[CumulativePD] ), FactPDTermStructure[HorizonMonths] = 12 )", "12-month PD by starting bucket.", "0.0%")
m(9, "PD", "Param Pool LGD", 'CALCULATE ( MAX ( DimParameters[Value] ), DimParameters[Parameter] = "Pool net LGD (severity)" )', "Pool LGD read from DimParameters (calibrated on dynamic_loss_monthly).", "0.0%")
m(9, "PD", "Param Base CPR", 'CALCULATE ( MAX ( DimParameters[Value] ), DimParameters[Parameter] = "Base annual CPR" )', "Base annual CPR read from DimParameters.", "0.0%")
m(9, "PD", "Param Rate Elasticity", 'CALCULATE ( MAX ( DimParameters[Value] ), DimParameters[Parameter] = "Rate shock PD elasticity (per 100bps)" )', "ASSUMPTION: PD uplift per 100bps rate shock.", "0.00")
m(9, "Challenger", "Challenger ECL (DAX)", f'''
VAR LGDp = [Param Pool LGD]
RETURN
SUMX (
    {T},
    VAR stg = {T}[IFRS9_Stage]
    VAR h   = IF ( stg = 1, 12, MIN ( MAX ( {T}[RemainingTerm], 1 ), 60 ) )
    VAR pdv = IF ( stg = 3, 1,
                LOOKUPVALUE ( FactPDTermStructure[CumulativePD],
                              FactPDTermStructure[DPD_Bucket], {T}[DelinquencyStatus],
                              FactPDTermStructure[HorizonMonths], h ) )
    VAR hh  = IF ( stg = 3, 6, h )
    VAR dfc = POWER ( 1 + RELATED ( DimLoan[InterestRate] ) / 100, - ( hh / 12 ) / 2 )
    RETURN IF ( {T}[EAD] <= 0, 0, pdv * LGDp * {T}[EAD] * dfc )
)''', "Independent IFRS 9 ECL: Stage 1 = 12M PD, Stage 2 = lifetime PD (capped 60M) from the roll-rate matrix, Stage 3 = 100%; LGD from the dynamic loss data; discounted at EIR.", "#,##0")
m(9, "Challenger", "Challenger vs Tape Variance %", "DIVIDE ( [Challenger ECL (DAX)] - [Total ECL (Tape)], [Total ECL (Tape)] )", "How far the data-driven ECL sits from the tape provision.", "0.0%")
m(9, "Staging", "Stage 2 DPD Value", "SELECTEDVALUE ( ParamStage2DPD[Value], 30 )", "Stage 2 DPD threshold (tape rule 30; brief 60).", "0")
m(9, "Staging", "Stage 3 DPD Value", "SELECTEDVALUE ( ParamStage3DPD[Value], 90 )", "Stage 3 DPD threshold (tape rule 90; brief 120).", "0")
m(9, "Staging", "ECL (Alt Staging)", f'''
VAR S2  = [Stage 2 DPD Value]
VAR S3  = [Stage 3 DPD Value]
VAR pd1 = CALCULATE ( DIVIDE ( SUMX ( {T}, {T}[PD_Estimate] * {T}[EAD] ), SUM ( {T}[EAD] ) ), REMOVEFILTERS (), {T}[IFRS9_Stage] = 1 )
VAR pd2 = CALCULATE ( DIVIDE ( SUMX ( {T}, {T}[PD_Estimate] * {T}[EAD] ), SUM ( {T}[EAD] ) ), REMOVEFILTERS (), {T}[IFRS9_Stage] = 2 )
VAR pd3 = CALCULATE ( DIVIDE ( SUMX ( {T}, {T}[PD_Estimate] * {T}[EAD] ), SUM ( {T}[EAD] ) ), REMOVEFILTERS (), {T}[IFRS9_Stage] = 3 )
RETURN
SUMX ( {T},
    VAR s = IF ( {T}[IsDefaulted] = 1 || {T}[DelinquencyDays] >= S3, 3, IF ( {T}[DelinquencyDays] >= S2, 2, 1 ) )
    RETURN SWITCH ( s, 1, pd1, 2, pd2, pd3 ) * {T}[LGD_Estimate] * {T}[EAD] )''', "Re-stages loans with adjustable DPD thresholds and applies stage-average PDs; compare the brief's 60/120 rule with the tape's 30/90 (approximation - use for impact, not as the provision).", "#,##0")

# =========================================================================== DAY 10 - STRESS
m(10, "What-if", "Default Rate Multiplier Value", "SELECTEDVALUE ( ParamDefaultMult[Value], 1 )", "PD multiplier slider.", "0.0")
m(10, "What-if", "Recovery Haircut Value", "SELECTEDVALUE ( ParamRecoveryHaircut[Value], 0 )", "Additive LGD increase slider (recovery haircut).", "0%")
m(10, "What-if", "CPR Change Value", "SELECTEDVALUE ( ParamCPRChange[Value], 0 )", "Change in annual CPR slider.", "0%")
m(10, "What-if", "Rate Shock Value", "SELECTEDVALUE ( ParamRateShock[Value], 0 )", "Interest-rate shock (bps) slider.", "0")
m(10, "What-if", "Effective PD Multiplier", "[Default Rate Multiplier Value] * ( 1 + [Param Rate Elasticity] * [Rate Shock Value] / 100 )", "Default multiplier compounded with the rate-shock PD elasticity (assumption).", "0.00")
m(10, "What-if", "Stress ECL (What-If)", f'''
VAR PDM = [Effective PD Multiplier]
VAR Hc  = [Recovery Haircut Value]
VAR cpr0 = [Param Base CPR]
VAR RunOff = POWER ( DIVIDE ( 1 - ( cpr0 + [CPR Change Value] ), 1 - cpr0 ), 0.5 )
RETURN
SUMX ( {T}, MIN ( 1, {T}[PD_Estimate] * PDM ) * MIN ( 1, {T}[LGD_Estimate] + Hc ) * {T}[EAD] * RunOff )''', "ECL under the four what-if levers (equals tape ECL when all levers are at their defaults). Faster prepayment shortens exposure (half-year average run-off).", "#,##0")
m(10, "What-if", "Stress ECL Uplift", "[Stress ECL (What-If)] - [Total ECL (Tape)]", "Increase vs tape provision.", "#,##0")
m(10, "What-if", "Stress ECL Uplift %", "DIVIDE ( [Stress ECL Uplift], [Total ECL (Tape)] )", "Percentage increase vs tape provision.", "0.0%")
m(10, "What-if", "Stress Coverage %", "DIVIDE ( [Stress ECL (What-If)], [Total EAD] )", "Stressed provision / EAD.", "0.00%")
m(10, "Scenarios", "Scenario ECL", f'''
VAR PDM = SELECTEDVALUE ( DimScenario[pd_multiplier], 1 )
VAR Hc  = SELECTEDVALUE ( DimScenario[lgd_add], 0 )
RETURN SUMX ( {T}, MIN ( 1, {T}[PD_Estimate] * PDM ) * MIN ( 1, {T}[LGD_Estimate] + Hc ) * {T}[EAD] )''', "ECL under the scenario in context (DimScenario on axis); rate-shock effect is embedded in the scenario PD multiplier.", "#,##0")
m(10, "Scenarios", "Probability-Weighted Scenario ECL", "SUMX ( ALL ( DimScenario ), DimScenario[weight] * [Scenario ECL] )", "IFRS 9 probability-weighted ECL across scenarios.", "#,##0")
m(10, "Scenarios", "Weighted ECL vs Tape", "[Probability-Weighted Scenario ECL] - [Total ECL (Tape)]", "Scenario-weighting overlay vs tape provision.", "#,##0")
m(10, "Scenarios", "Cum Loss % at Horizon", "VAR p = CALCULATE ( MAX ( FactPool[period] ), ALL ( FactPool ) )\nRETURN CALCULATE ( MAX ( FactPool[cumulative_loss_pct] ), FactPool[period] = p )", "Cumulative projected net loss at the end of the 36-month horizon.", "0.00%")
m(10, "Scenarios", "Senior Tranche Loss", 'CALCULATE ( [Tranche Writedown], DimTranche[tranche] = "Class A (Senior)" )', "Loss absorbed by Class A.", "#,##0")
m(10, "Scenarios", "Mezzanine Tranche Loss", 'CALCULATE ( [Tranche Writedown], DimTranche[tranche] = "Class B (Mezzanine)" )', "Loss absorbed by Class B.", "#,##0")
m(10, "Scenarios", "Junior Tranche Loss %", 'CALCULATE ( [Writedown % of Tranche], DimTranche[tranche] = "Class C (Junior)" )', "Class C write-down %.", "0.0%")
m(10, "Scenarios", "Equity Loss %", 'CALCULATE ( [Writedown % of Tranche], DimTranche[tranche] = "Equity / Residual" )', "Equity write-down %.", "0.0%")
m(10, "Scenarios", "Minimum OC Ratio", "CALCULATE ( MIN ( FactPool[oc_ratio] ) )", "Lowest OC ratio over the horizon.", "0.000")
m(10, "Scenarios", "CNL Trigger Headroom", "[CNL Trigger Value] - [Cum Loss % at Horizon]", "Distance to the CNL trigger at horizon (negative = breached).", "0.00%")
m(10, "Scenarios", "OC Trigger Headroom", "[Minimum OC Ratio] - [OC Trigger Value]", "Distance to the OC trigger at the tightest point.", "0.000")

# ---- Risk Arena (gamified scoring; weights live in DimParameters, none hard-coded)
for lab in ["Equity", "Junior", "Mezzanine", "Senior", "Trigger breach"]:
    m(10, "Arena", f"Arena W {lab}", f'CALCULATE ( MAX ( DimParameters[Value] ), DimParameters[Parameter] = "Arena weight - {lab}" )', f"Points lost per full {lab.lower()} loss / breach (DimParameters).", "0")
for lab in ["Fortress", "Resilient", "Stressed"]:
    m(10, "Arena", f"Arena Min {lab}", f'CALCULATE ( MAX ( DimParameters[Value] ), DimParameters[Parameter] = "Arena rating - {lab} (min score)" )', f"Minimum score for the {lab} rating (DimParameters).", "0")
m(10, "Arena", "Mezzanine Loss %", 'CALCULATE ( [Writedown % of Tranche], DimTranche[tranche] = "Class B (Mezzanine)" )', "Class B write-down %.", "0.0%")
m(10, "Arena", "Senior Loss %", 'CALCULATE ( [Writedown % of Tranche], DimTranche[tranche] = "Class A (Senior)" )', "Class A write-down %.", "0.0%")
m(10, "Arena", "Arena Score", """
MAX ( 0, ROUND ( 100
    - [Arena W Equity]         * [Equity Loss %]
    - [Arena W Junior]         * [Junior Tranche Loss %]
    - [Arena W Mezzanine]      * [Mezzanine Loss %]
    - [Arena W Senior]         * [Senior Loss %]
    - [Arena W Trigger breach] * [Trigger Breach Count], 0 ) )""", "Crisis-survival score 0-100 for the scenario in context: starts at 100, loses points for write-downs by seniority and for each trigger breached.", "0")
m(10, "Arena", "Arena Rating", """
SWITCH ( TRUE (),
    [Arena Score] >= [Arena Min Fortress],  "Fortress",
    [Arena Score] >= [Arena Min Resilient], "Resilient",
    [Arena Score] >= [Arena Min Stressed],  "Stressed",
    "Crisis" )""", "Rating band from the score.", "")
m(10, "Arena", "Arena Message", """
"Scenario: " & SELECTEDVALUE ( DimScenario[scenario], "All" )
    & " | Cum. loss " & FORMAT ( [Cum Loss % at Horizon], "0.0%" )
    & " | Score " & [Arena Score] & " - " & [Arena Rating]""", "Narrative line for the Arena scoreboard card.", "")

# =========================================================================== DAY 11 - INVESTOR REPORTING
m(11, "Loss & recovery", "Cumulative Net Loss (Reporting)", f"VAR LastD = MAX ( {Y}[MonthEnd] )\nRETURN CALCULATE ( [Net Loss (Monthly)], FILTER ( ALL ( DimDate[Date] ), DimDate[Date] <= LastD ) )", "Cumulative net loss over the reporting history.", "#,##0")
m(11, "Loss & recovery", "Cumulative Net Loss % of Original", "DIVIDE ( [Cumulative Net Loss (Reporting)], CALCULATE ( [Original Balance], REMOVEFILTERS () ) )", "Cumulative net loss / original pool balance.", "0.00%")
m(11, "Loss & recovery", "Annualised Default Rate (All Months)", f"1 - POWER ( 1 - DIVIDE ( SUM ( {Y}[NewDefaults_Balance] ), SUM ( {Y}[BOP_Balance] ) ), 12 )", "Annualised default rate using all 12 monthly rows (used to calibrate the base CDR).", "0.00%")
m(11, "Prepayment", "SMM", f"DIVIDE ( SUM ( {Y}[Prepayments_ThisMonth] ), SUM ( {Y}[BOP_Balance] ) )", "Single monthly mortality.", "0.000%")
m(11, "Prepayment", "CPR (Annualised)", "1 - POWER ( 1 - [SMM], 12 )", "Annualised conditional prepayment rate.", "0.0%")
m(11, "Collections", "Collection Efficiency", f"DIVIDE ( SUM ( {Y}[CollectionsTotal] ), SUM ( {Y}[BillingAmount] ) )", "Collections / billing.", "0.0%")
m(11, "Collections", "Excess Spread (Reported)", f"SUM ( {Y}[ExcessSpread_Monthly] )", "Excess spread as reported in the servicer file.", "#,##0")
m(11, "Collections", "Excess Spread % (Reported, annualised)", f"DIVIDE ( SUM ( {Y}[ExcessSpread_Monthly] ), SUM ( {Y}[BOP_Balance] ) ) * 12", "Annualised reported excess spread / opening balance.", "0.00%")
m(11, "Report header", "Report Month Label", f'FORMAT ( MAX ( {Y}[MonthEnd] ), "mmm yyyy" )', "Header label for servicer-report pages.", "")
m(11, "Data tape", "Data Tape Field Completeness %", 'DIVIDE ( COUNTROWS ( FILTER ( DataFieldMapping, DataFieldMapping[Status] IN { "Present", "Computed" } ) ), COUNTROWS ( DataFieldMapping ) )', "Share of rating-agency fields (Brief A2.1) present or computable from the supplied data.", "0.0%")

# =========================================================================== DAY 12 - GOVERNANCE
m(12, "Security", "Current User", "USERPRINCIPALNAME ()", "Signed-in user (for RLS testing).", "")
m(12, "Security", "Current User Role", 'LOOKUPVALUE ( DimSecurity[Role], DimSecurity[UserEmail], USERPRINCIPALNAME (), "Unmapped" )', "Role of the signed-in user per DimSecurity.", "")
m(12, "Security", "Loans Visible to User", "[Loan Count]", "Row-count check under 'View as role' (must differ by role).", "#,##0")
m(12, "Validation", "Reconciliation Break (Tape vs Dynamic)", f'''
VAR LastSeq = CALCULATE ( MAX ( {Y}[ReportingSeq] ), REMOVEFILTERS () )
RETURN CALCULATE ( [Total Outstanding Balance], REMOVEFILTERS () )
     - CALCULATE ( SUM ( {Y}[EOP_Balance] ), REMOVEFILTERS (), {Y}[ReportingSeq] = LastSeq )''', "Tape balance minus the dynamic file's closing balance (reconciliation break, DQ10).", "#,##0")
m(12, "Validation", "High-Severity Data Issues", 'CALCULATE ( COUNTROWS ( DataQualityLog ), DataQualityLog[severity] = "High" )', "Number of high-severity data-quality findings.", "0")
m(12, "Validation", "ECL Check Flag", 'IF ( ABS ( [ECL Recompute Variance] ) < 1, "PASS", "FAIL" )', "Pass/fail flag for the PD x LGD x EAD recomputation.", "")

# --------------------------------------------------------------------------- dedupe sanity
names = [x[3] for x in M]
assert len(names) == len(set(names)), [n for n in names if names.count(n) > 1]
DAYS = {3: "Core pool, delinquency and weighted-average measures", 4: "Time intelligence and trend alerts", 5: "Advanced DAX, virtual tables, concentration",
        6: "Waterfall, tranche allocation and trigger testing", 7: "DPD distribution, roll rates and ageing", 8: "Vintage / static pool analysis",
        9: "IFRS 9 ECL framework", 10: "Stress testing and what-if", 11: "Investor reporting and data tape", 12: "Security, governance and validation"}


def deps(dax):
    return ", ".join(sorted({x for x in re.findall(r"\[([^\]]+)\]", dax) if x in names}))


EXP = pd.read_csv("powerbi_data/ExpectedMeasureValues.csv").set_index("Measure")["ExpectedValue_Python"].to_dict()
missing = [k for k in EXP if k not in names]
assert not missing, f"expected-value names without measure: {missing}"

# =========================================================================== MARKDOWN
L = ["# Securitisation Risk Analytics - DAX Measure Library", "",
     f"**{len(M)} measures** across Days 3-12 of the brief. Every measure is formula-driven (no hard-coded results). "
     "Generated from `build_measures.py`; the same list feeds the Excel measure dictionary and the DAX Studio benchmark file.", "",
     "## 1. Model and relationships", "",
     "Load the CSVs from `powerbi_data/`. Create a blank table `_Measures` and paste measures into it.", "",
     "| From (one side) | To (many side) | Notes |", "|---|---|---|",
     "| DimLoan[LoanID] | FactLoanTape[LoanID] | one-to-many, single direction |",
     "| DimLoan[LoanID] | FactMonthlyPerformance[LoanID] | one-to-many, single |",
     "| DimBorrower[BorrowerID] | DimLoan[BorrowerID] | one-to-many, single (borrower IDs are unique) |",
     "| DimVehicle[LoanID] | DimLoan[LoanID] | one-to-many, single |",
     "| DimGeography[State] | DimLoan[State] | one-to-many, single |",
     "| DimVintage[VintageID] | DimLoan[VintageID] | one-to-many, single |",
     "| DimVintage[VintageID] | FactStaticPool[VintageID] | one-to-many, single |",
     "| DimStage[IFRS9_Stage] | FactLoanTape[IFRS9_Stage] | one-to-many, single |",
     "| DimBucket[DPD_Bucket] | FactMonthlyPerformance[DPD_Bucket] | one-to-many, single |",
     "| DimDate[Date] | FactMonthlyPerformance[MonthEnd] | one-to-many, single |",
     "| DimDate[Date] | FactDynamicLoss[MonthEnd] | one-to-many, single |",
     "| DimDate[Date] | DimLoan[OriginationDate] | **inactive** role-playing date (use USERELATIONSHIP) |",
     "| DimScenario[scenario_id] | FactPool[scenario_id], FactWaterfall[scenario_id] | one-to-many, single |",
     "| DimPeriod[period] | FactPool[period], FactWaterfall[period] | one-to-many, single |",
     "| DimTranche[tranche] | FactWaterfall[tranche] | one-to-many, single |", "",
     "**Disconnected tables (no relationships):** FactRollRate, FactPDTermStructure, DimParameters, DimSecurity, DataFieldMapping, DataQualityLog, ExpectedMeasureValues.",
     "**Do not** relate DimBucket to FactLoanTape or DimDate to FactLoanTape: together with DimLoan they would create an ambiguous path. Slice the tape by `DelinquencyStatus` (sorted by `BucketOrder`).", "",
     "Mark DimDate as the date table (column Date). Sort DelinquencyStatus / DPD_Bucket by BucketOrder, SnapshotLabel by SnapshotSeq, scenario by scenario_id, tranche by seniority.", "",
     "## 2. Parameter tables (Modeling > New table)", "", "```dax"]
for nm, expr, desc in PARAM_TABLES:
    L.append(f"{nm} = {expr}    // {desc}")
L += ["```", "", "Each table is a single column named `Value`; attach a slicer to it. The matching `... Value` measures below read the selection with SELECTEDVALUE and a default.", ""]
sec = 3
for day in sorted(DAYS):
    L += [f"## {sec}. Day {day} - {DAYS[day]}", ""]
    sec += 1
    cat = None
    for (mid, d, c, name, dax, purpose, fmt) in M:
        if d != day:
            continue
        L += [f"### {mid} - {name}", f"*{c}. {purpose}*  " + (f"Format: `{fmt}`" if fmt else ""), "", "```dax", f"{name} =", dax if "\n" in dax else dax, "```", ""]
L += ["## Notes and assumptions", "",
      "- **Trigger levels** (12% 60+ DPD, 2% CNL, 1.05 OC), tranche structure, coupons and the 0.5% servicing fee are **assumptions** - the brief supplies none.",
      "- **Duplicate 30-May snapshot**: roll-rate and PD measures use all 12 sequential snapshots; date-based trend measures use IncludeInTrend = 1.",
      "- **CNL Rate @ MOB / Cumulative Default Rate @ MOB** are cumulative series: always put MonthsOnBook on the axis (or filter to one MOB) to avoid summing cumulative values across MOBs.",
      "- **Challenger vs tape ECL** differ materially because PD in the tape (about 2% Stage 1) is far below the roll-rate-implied PD; this is reported as a finding, not corrected.",
      "- Measures could not be run here (no Power BI engine). Use `ExpectedMeasureValues.csv` / the dictionary's *Expected value* column and DAX Studio to confirm each headline figure."]
Path("securitisation_measures.dax.md").write_text("\n".join(L), encoding="utf-8")

# =========================================================================== DAX STUDIO FILE
B = ["// DAX Studio benchmark - run with Server Timings ON; target < 100 ms per measure (brief Day 12).",
     "// One EVALUATE per measure. Paste into DAX Studio connected to the open .pbix, or use 'Run Benchmark'.", ""]
for (mid, d, c, name, dax, purpose, fmt) in M:
    B.append(f'// {mid}\nEVALUATE ROW ( "Measure", "{name}", "Value", [{name}] )\n')
Path("DAXStudio_Benchmark.dax").write_text("\n".join(B), encoding="utf-8")

# =========================================================================== EXCEL DICTIONARY
wb = Workbook()
ws = wb.active
ws.title = "Measure Dictionary"
hdr = ["ID", "Day", "Category", "Measure", "DAX formula", "Business purpose", "Depends on", "Format", "Expected value (Python check)",
       "Value in Power BI", "Matches? (formula)", "DAX Studio ms", "Tested (Y/N)"]
ws.append(hdr)
for (mid, d, c, name, dax, purpose, fmt) in M:
    r = ws.max_row + 1
    exp = EXP.get(name)
    ws.append([mid, d, c, name, dax, purpose, deps(dax), fmt, exp, None,
               f'=IF(OR(J{r}="",I{r}=""),"",IF(ABS(J{r}-I{r})<=MAX(0.0001,ABS(I{r})*0.0001),"MATCH","CHECK"))' if exp is not None else "", None, None])
F = "Arial"
for c in ws[1]:
    c.font, c.fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B2545")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(vertical="top", wrap_text=c.column in (5, 6, 7))
    row[8].font = Font(name=F, size=10, color="0000FF")
    row[9].fill = PatternFill("solid", fgColor="FFFF99")
for col, w in zip("ABCDEFGHIJKLM", [8, 5, 16, 32, 70, 48, 34, 9, 18, 16, 14, 12, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:M{ws.max_row}"
r0 = ws.max_row + 2
ws.cell(r0, 4, "Legend").font = Font(name=F, bold=True)
ws.cell(r0 + 1, 4, "Blue = value computed independently in Python (pandas) from the supplied CSVs.").font = Font(name=F, size=9, color="0000FF")
ws.cell(r0 + 2, 4, "Yellow = paste the value your Power BI measure returns (card, no filters); column K then shows MATCH / CHECK.").font = Font(name=F, size=9)

w2 = wb.create_sheet("Summary")
w2.append(["Day", "Focus", "Measures"])
for d in sorted(DAYS):
    w2.append([d, DAYS[d], f"=COUNTIF('Measure Dictionary'!B:B,{d})"])
w2.append(["", "Total", f"=SUM(C2:C{w2.max_row})"])
w2.append(["", "Brief requirement", 80])
w2.append(["", "Meets 80+ requirement?", f'=IF(C{w2.max_row - 1}>=C{w2.max_row},"YES","NO")'])
for c in w2[1]:
    c.font, c.fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B2545")
for row in w2.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10)
w2.cell(w2.max_row - 1, 3).font = Font(name=F, size=10, color="0000FF")
w2.column_dimensions["B"].width, w2.column_dimensions["C"].width = 58, 12

w3 = wb.create_sheet("Parameters")
w3.append(["Parameter table", "DAX", "Purpose"])
for nm, expr, desc in PARAM_TABLES:
    w3.append([nm, expr, desc])
for c in w3[1]:
    c.font, c.fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B2545")
for row in w3.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10)
for col, w in zip("ABC", [24, 40, 55]):
    w3.column_dimensions[col].width = w
wb.save("483553A_ManyaSyal_DAX_Measure_Dictionary.xlsx")
print(f"{len(M)} measures | by day:", {d: sum(1 for x in M if x[1] == d) for d in sorted(DAYS)})
