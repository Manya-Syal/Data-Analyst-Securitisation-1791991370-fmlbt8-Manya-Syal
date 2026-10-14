"""
Builds 483553A_ManyaSyal_ECL_Waterfall_Validation.xlsx  (Brief Day 9 'Excel crosscheck' + Day 12 'waterfall validation').
All calculations are live Excel formulas; hard-coded numbers are blue and sourced from powerbi_data/*.csv.
Run after securitisation_pipeline.py.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

D = "powerbi_data/"
tape = pd.read_csv(D + "FactLoanTape.csv")
loan = pd.read_csv(D + "DimLoan.csv")[["LoanID", "InterestRate"]]
tape = tape.merge(loan, on="LoanID")
pdts = pd.read_csv(D + "FactPDTermStructure.csv")
par = pd.read_csv(D + "DimParameters.csv").set_index("Parameter")["Value"]
sc = pd.read_csv(D + "DimScenario.csv")
fp = pd.read_csv(D + "FactPool.csv")
fw = pd.read_csv(D + "FactWaterfall.csv")
tr = pd.read_csv(D + "DimTranche.csv").sort_values("seniority")

F = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HDR = PatternFill("solid", fgColor="0B2545")
YEL = PatternFill("solid", fgColor="FFFF00")
f_in = Font(name=F, size=10, color=BLUE)
f_n = Font(name=F, size=10, color=BLACK)
f_b = Font(name=F, size=10, bold=True)
f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_t = Font(name=F, size=14, bold=True, color="0B2545")


def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill = f_h, HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def put(ws, ref, v, font=f_n, fmt=None):
    ws[ref] = v
    ws[ref].font = font
    if fmt:
        ws[ref].number_format = fmt


wb = Workbook()

# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
put(ws, "A1", "Securitisation Risk Analytics - Excel Validation Workbook", f_t)
lines = [
    "Purpose: independent cross-check of (1) IFRS 9 ECL and (2) the 36-month tranche waterfall used in the Power BI model.",
    "Colour code: BLUE = hard-coded input or value pasted from the Python / CSV outputs; BLACK = formula. Yellow cell = the one you change.",
    "",
    "Sheets",
    "  Assumptions       - calibrated parameters (from dynamic_loss_monthly and the tape), scenario table, tranche structure. Change the scenario in B10.",
    "  Waterfall         - month-by-month waterfall rebuilt in Excel formulas; compares to the Python engine and checks cash conservation.",
    "  Python_Results    - end balances produced by securitisation_pipeline.py for all six scenarios (comparison source).",
    "  ECL_LoanLevel     - loan-level recomputation: tape ECL vs PD x LGD x EAD, and the challenger ECL from the transition-matrix PD term structure.",
    "  PD_TermStructure  - cumulative PD by starting DPD bucket and horizon (from the empirical roll-rate matrix).",
    "  ECL_Summary       - stage-level reconciliation and PASS/FAIL checks.",
    "",
    "Key assumptions (also listed in Assumptions): tranche sizes/coupons, 0.5% servicing fee and trigger levels are ASSUMED - the brief gives none.",
    "Challenger LGD = net loss severity from dynamic_loss_monthly (NetLoss / NewDefaults_Balance).",
    "Discounting: ECL discounted at the loan EIR over half the PD horizon (average timing of default). Stage 3 = PD 100%, 6-month horizon.",
    "Data issues found in the supplied files are in powerbi_data/DataQualityLog.csv - read them before relying on any figure.",
]
for i, t in enumerate(lines, start=3):
    put(ws, f"A{i}", t, f_b if t == "Sheets" else f_n)
ws.column_dimensions["A"].width = 150

# ------------------------------------------------------------------ Assumptions
A = wb.create_sheet("Assumptions")
put(A, "A1", "Assumptions and scenario selector", f_t)
rows = [
    (3, "Pool net LGD (severity)", par["Pool net LGD (severity)"], "0.00%", "dynamic_loss_monthly: NetLoss / NewDefaults_Balance"),
    (4, "Base annual CDR", par["Base annual CDR"], "0.00%", "dynamic_loss_monthly: NewDefaults_Balance / BOP_Balance, annualised"),
    (5, "Base annual CPR", par["Base annual CPR"], "0.00%", "dynamic_loss_monthly: Prepayments / BOP_Balance, annualised"),
    (6, "Pool WAC", par["Pool WAC"], "0.00%", "tape: InterestRate weighted by CurrentBalance"),
    (7, "Pool WAM (months)", par["Pool WAM (months)"], "0.0", "tape: RemainingTerm weighted by CurrentBalance"),
    (8, "Pool balance", par["Pool balance"], "#,##0", "tape: sum CurrentBalance"),
    (9, "Servicing fee (p.a.)", 0.005, "0.00%", "ASSUMPTION"),
]
for r, lab, v, fmt, src in rows:
    put(A, f"A{r}", lab)
    put(A, f"B{r}", float(v), f_in, fmt)
    put(A, f"C{r}", src, Font(name=F, size=9, italic=True))
put(A, "A10", "SELECTED SCENARIO ID (1-6)", f_b)
put(A, "B10", 4, f_in)
A["B10"].fill = YEL
dv = DataValidation(type="list", formula1='"1,2,3,4,5,6"', allow_blank=False)
A.add_data_validation(dv)
dv.add("B10")
put(A, "C10", "Change this cell; the Waterfall sheet and its Python comparison update. 1=Base ... 6=Rate Shock", Font(name=F, size=9, italic=True))
S0 = 18   # first scenario row
labs = [(11, "Selected annual CDR", f"=INDEX($F${S0}:$F${S0 + 5},MATCH($B$10,$A${S0}:$A${S0 + 5},0))", "0.000%"),
        (12, "Selected annual CPR", f"=INDEX($G${S0}:$G${S0 + 5},MATCH($B$10,$A${S0}:$A${S0 + 5},0))", "0.000%"),
        (13, "Selected pool LGD", f"=INDEX($H${S0}:$H${S0 + 5},MATCH($B$10,$A${S0}:$A${S0 + 5},0))", "0.00%"),
        (14, "Monthly default rate (MDR)", "=1-(1-B11)^(1/12)", "0.0000%"),
        (15, "Single monthly mortality (SMM)", "=1-(1-B12)^(1/12)", "0.0000%")]
for r, lab, f, fmt in labs:
    put(A, f"A{r}", lab)
    put(A, f"B{r}", f, f_n, fmt)
hdr(A, S0 - 1, ["ID", "Scenario", "Weight", "CDR multiple", "CPR change", "Annual CDR", "Annual CPR", "Pool LGD", "PD multiplier", "LGD add"])
for i, s in enumerate(sc.itertuples()):
    r = S0 + i
    put(A, f"A{r}", int(s.scenario_id), f_in)
    put(A, f"B{r}", s.scenario, f_in)
    put(A, f"C{r}", float(s.weight), f_in, "0%")
    put(A, f"D{r}", float(s.cdr_multiple), f_in, "0.0")
    put(A, f"E{r}", float(s.cpr_change), f_in, "0.0%")
    put(A, f"F{r}", f"=ROUND(MIN(0.6,$B$4*D{r}),5)", f_n, "0.000%")
    put(A, f"G{r}", f"=ROUND(MAX(0.01,$B$5+E{r}),5)", f_n, "0.000%")
    put(A, f"H{r}", f"=ROUND(MIN(0.95,MAX(0.1,$B$3+J{r})),4)", f_n, "0.00%")
    put(A, f"I{r}", float(s.pd_multiplier), f_in, "0.0")
    put(A, f"J{r}", float(s.lgd_add), f_in, "0.0%")
T0 = 27
hdr(A, T0 - 1, ["Tranche", "Size % of pool", "Coupon", "Attachment", "Detachment", "Note"])
for i, t in enumerate(tr.itertuples()):
    r = T0 + i
    put(A, f"A{r}", t.tranche, f_in)
    put(A, f"B{r}", float(t.size_pct), f_in, "0.0%")
    put(A, f"C{r}", float(t.coupon), f_in, "0.0%")
    put(A, f"D{r}", float(t.attachment), f_in, "0.0%")
    put(A, f"E{r}", float(t.detachment), f_in, "0.0%")
    put(A, f"F{r}", "ASSUMED structure - brief supplies no tranche terms", Font(name=F, size=9, italic=True))
put(A, f"A{T0 + 4}", "Check: sizes sum to 100%")
put(A, f"B{T0 + 4}", f"=SUM(B{T0}:B{T0 + 3})", f_n, "0.0%")
for col, w in zip("ABCDEFGHIJ", [34, 18, 70, 14, 14, 12, 12, 12, 14, 10]):
    A.column_dimensions[col].width = w
szA, szB, szC, szE = [f"Assumptions!$B${T0 + i}" for i in range(4)]
cA, cB, cC = [f"Assumptions!$C${T0 + i}" for i in range(3)]

# ------------------------------------------------------------------ Python_Results
PR = wb.create_sheet("Python_Results")
hdr(PR, 1, ["scenario_id", "period", "ending_pool", "Class A end", "Class B end", "Class C end", "Equity end"])
piv = fw.pivot_table(index=["scenario_id", "period"], columns="tranche", values="ending_balance", aggfunc="sum").reset_index()
mg = fp[["scenario_id", "period", "ending_pool"]].merge(piv, on=["scenario_id", "period"])
cols = ["scenario_id", "period", "ending_pool", "Class A (Senior)", "Class B (Mezzanine)", "Class C (Junior)", "Equity / Residual"]
for i, row in enumerate(mg[cols].itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        c = PR.cell(i, j, float(v) if j > 2 else int(v))
        c.font = f_in
        if j > 2:
            c.number_format = "#,##0.00"
for col in "ABCDEFG":
    PR.column_dimensions[col].width = 16
PRN = len(mg) + 1

# ------------------------------------------------------------------ Waterfall
W = wb.create_sheet("Waterfall")
put(W, "A1", "Waterfall validation - Excel rebuild vs Python engine", f_t)
put(W, "A2", "Selected scenario:")
put(W, "B2", "=INDEX(Assumptions!$B$18:$B$23,MATCH(Assumptions!$B$10,Assumptions!$A$18:$A$23,0))", f_b)
put(W, "D2", "Max abs difference vs Python (any column):")
put(W, "H2", "=MAX(MAX(AW7:BA42),-MIN(AW7:BA42))", f_b, "#,##0.0000")
put(W, "I2", '=IF(H2<1,"PASS","FAIL")', f_b)
put(W, "D3", "Max abs cash-conservation gap (notes - pool):")
put(W, "H3", "=MAX(MAX(AQ7:AQ42),-MIN(AQ7:AQ42))", f_b, "#,##0.0000")
put(W, "I3", '=IF(H3<1,"PASS","FAIL")', f_b)
heads = ["Period", "Beg pool", "Defaults", "Performing", "Prepay", "Sched principal", "Interest coll", "Recoveries", "Net loss", "Servicing fee",
         "Avail interest", "Beg A", "Beg B", "Beg C", "Beg Equity", "Coupon due A", "Coupon due B", "Coupon due C", "Paid A", "Paid B", "Paid C",
         "Excess spread", "Cover for losses", "Residual to equity", "WD Equity", "WD C", "WD B", "WD A", "A after WD", "C after WD", "B after WD",
         "Eq after WD", "Principal avail", "Prin A", "Prin B", "Prin C", "Prin Eq", "End A", "End B", "End C", "End Equity", "End pool", "Notes - pool",
         "Python end pool", "Python A", "Python B", "Python C", "Python Eq", "Diff pool", "Diff A", "Diff B", "Diff C", "Diff Eq"]
# build header positions: columns A..AP are calc, then AQ check, AR..AV python, AW..BA diffs
order = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z",
         "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU", "AV", "AW", "AX", "AY", "AZ", "BA"]
for lab, col in zip(heads, order):
    c = W[f"{col}6"]
    c.value, c.font, c.fill = lab, f_h, HDR
    c.alignment = Alignment(horizontal="center", wrap_text=True)
W.row_dimensions[6].height = 32
mdr, smm, lgd = "Assumptions!$B$14", "Assumptions!$B$15", "Assumptions!$B$13"
wac, wam, fee, pool0 = "Assumptions!$B$6", "Assumptions!$B$7", "Assumptions!$B$9", "Assumptions!$B$8"
for t in range(1, 37):
    r = 6 + t
    p = r - 1
    f = {}
    f["A"] = t
    f["B"] = f"={pool0}" if t == 1 else f"=AP{p}"
    f["C"] = f"=B{r}*{mdr}"
    f["D"] = f"=B{r}-C{r}"
    f["E"] = f"=D{r}*{smm}"
    f["F"] = f"=(D{r}-E{r})/MAX({wam}-A{r}+1,1)"
    f["G"] = f"=D{r}*{wac}/12"
    f["H"] = f"=C{r}*(1-{lgd})"
    f["I"] = f"=C{r}*{lgd}"
    f["J"] = f"=B{r}*{fee}/12"
    f["K"] = f"=MAX(G{r}-J{r},0)"
    f["L"] = f"={pool0}*{szA}" if t == 1 else f"=AL{p}"
    f["M"] = f"={pool0}*{szB}" if t == 1 else f"=AM{p}"
    f["N"] = f"={pool0}*{szC}" if t == 1 else f"=AN{p}"
    f["O"] = f"={pool0}*{szE}" if t == 1 else f"=AO{p}"
    f["P"] = f"=L{r}*{cA}/12"
    f["Q"] = f"=M{r}*{cB}/12"
    f["R"] = f"=N{r}*{cC}/12"
    f["S"] = f"=MIN(P{r},K{r})"
    f["T"] = f"=MIN(Q{r},K{r}-S{r})"
    f["U"] = f"=MIN(R{r},K{r}-S{r}-T{r})"
    f["V"] = f"=K{r}-S{r}-T{r}-U{r}"
    f["W"] = f"=MIN(V{r},I{r})"
    f["X"] = f"=V{r}-W{r}"
    f["Y"] = f"=MIN(O{r},I{r}-W{r})"
    f["Z"] = f"=MIN(N{r},I{r}-W{r}-Y{r})"
    f["AA"] = f"=MIN(M{r},I{r}-W{r}-Y{r}-Z{r})"
    f["AB"] = f"=MIN(L{r},I{r}-W{r}-Y{r}-Z{r}-AA{r})"
    f["AC"] = f"=L{r}-AB{r}"
    f["AD"] = f"=N{r}-Z{r}"
    f["AE"] = f"=M{r}-AA{r}"
    f["AF"] = f"=O{r}-Y{r}"
    f["AG"] = f"=F{r}+E{r}+H{r}+W{r}"
    f["AH"] = f"=MIN(AC{r},AG{r})"
    f["AI"] = f"=MIN(AE{r},AG{r}-AH{r})"
    f["AJ"] = f"=MIN(AD{r},AG{r}-AH{r}-AI{r})"
    f["AK"] = f"=MIN(AF{r},AG{r}-AH{r}-AI{r}-AJ{r})"
    f["AL"] = f"=AC{r}-AH{r}"
    f["AM"] = f"=AE{r}-AI{r}"
    f["AN"] = f"=AD{r}-AJ{r}"
    f["AO"] = f"=AF{r}-AK{r}"
    f["AP"] = f"=D{r}-E{r}-F{r}"
    f["AQ"] = f"=SUM(AL{r}:AO{r})-AP{r}"
    sel = f"Assumptions!$B$10"
    for col, src in zip(["AR", "AS", "AT", "AU", "AV"], "CDEFG"):
        f[col] = f"=SUMIFS(Python_Results!${src}$2:${src}${PRN},Python_Results!$A$2:$A${PRN},{sel},Python_Results!$B$2:$B${PRN},A{r})"
    f["AW"] = f"=AP{r}-AR{r}"
    f["AX"] = f"=AL{r}-AS{r}"
    f["AY"] = f"=AM{r}-AT{r}"
    f["AZ"] = f"=AN{r}-AU{r}"
    f["BA"] = f"=AO{r}-AV{r}"
    for col, v in f.items():
        c = W[f"{col}{r}"]
        c.value = v
        c.font = f_n if col != "A" else f_n
        c.number_format = "0" if col == "A" else "#,##0;(#,##0);-"
for col in order:
    W.column_dimensions[col].width = 14
W.freeze_panes = "B7"
put(W, "A44", "Logic: interest waterfall pays fees > Class A > B > C coupons; remaining excess spread first cures losses, the rest goes to equity. "
              "Uncovered losses write down Equity > C > B > A. Principal (scheduled + prepay + recoveries + excess used for losses) pays A > B > C > Equity sequentially.",
    Font(name=F, size=9, italic=True))

# ------------------------------------------------------------------ PD_TermStructure
P = wb.create_sheet("PD_TermStructure")
hdr(P, 1, ["DPD_Bucket", "HorizonMonths", "CumulativePD", "Key"])
for i, r in enumerate(pdts.itertuples(index=False), start=2):
    put(P, f"A{i}", r.DPD_Bucket, f_in)
    put(P, f"B{i}", int(r.HorizonMonths), f_in)
    put(P, f"C{i}", float(r.CumulativePD), f_in, "0.0000%")
    put(P, f"D{i}", f'=A{i}&"|"&B{i}', f_n)
for col, w in zip("ABCD", [16, 14, 14, 20]):
    P.column_dimensions[col].width = w
PN = len(pdts) + 1

# ------------------------------------------------------------------ ECL_LoanLevel
E = wb.create_sheet("ECL_LoanLevel")
put(E, "A1", "Loan-level ECL cross-check (tape vs recomputed vs challenger)", f_t)
put(E, "A2", "Blue = pasted from the tape / Python. Challenger columns L:Q are live formulas using PD_TermStructure and Assumptions!B3 (pool LGD).", Font(name=F, size=9, italic=True))
heads = ["LoanID", "IFRS9 stage", "DPD status", "Remaining term", "Interest rate %", "PD (tape)", "LGD (tape)", "EAD", "ECL (tape)",
         "ECL recomputed", "Variance", "Horizon (m)", "Challenger PD", "Discount factor", "Challenger ECL", "Python challenger ECL", "Diff vs Python"]
hdr(E, 5, heads)
n = len(tape)
for i, r in enumerate(tape.itertuples(index=False)):
    x = 6 + i
    vals = [r.LoanID, int(r.IFRS9_Stage), r.DelinquencyStatus, int(r.RemainingTerm), float(r.InterestRate), float(r.PD_Estimate), float(r.LGD_Estimate), float(r.EAD), float(r.ECL_Provision)]
    for j, v in enumerate(vals, start=1):
        c = E.cell(x, j, v)
        c.font = f_in
        c.number_format = "#,##0.00" if j >= 8 else ("0.0000" if j in (6, 7) else "General")
    E[f"J{x}"] = f"=F{x}*G{x}*H{x}"
    E[f"K{x}"] = f"=J{x}-I{x}"
    E[f"L{x}"] = f"=IF(B{x}=1,12,IF(B{x}=2,MIN(MAX(D{x},1),60),6))"
    E[f"M{x}"] = f'=IF(B{x}=3,1,IFERROR(INDEX(PD_TermStructure!$C$2:$C${PN},MATCH(C{x}&"|"&L{x},PD_TermStructure!$D$2:$D${PN},0)),0))'
    E[f"N{x}"] = f"=(1+E{x}/100)^(-(L{x}/12)/2)"
    E[f"O{x}"] = f"=IF(H{x}<=0,0,M{x}*Assumptions!$B$3*H{x}*N{x})"
    E[f"P{x}"] = float(r.ECL_Challenger)
    E[f"Q{x}"] = f"=O{x}-P{x}"
    for col in "JKLMNOPQ":
        E[f"{col}{x}"].font = f_in if col == "P" else f_n
    for col, fmt in zip("JKLMNOPQ", ["#,##0.00", "#,##0.00", "0", "0.00%", "0.0000", "#,##0.00", "#,##0.00", "#,##0.00"]):
        E[f"{col}{x}"].number_format = fmt
last = 5 + n
for col, w in zip("ABCDEFGHIJKLMNOPQ", [22, 8, 14, 10, 10, 10, 10, 14, 14, 14, 11, 10, 11, 10, 14, 16, 12]):
    E.column_dimensions[col].width = w
E.freeze_panes = "B6"

# ------------------------------------------------------------------ ECL_Summary
Z = wb.create_sheet("ECL_Summary")
put(Z, "A1", "ECL reconciliation by IFRS 9 stage", f_t)
hdr(Z, 3, ["Stage", "Loans", "EAD", "ECL (tape)", "ECL recomputed", "Variance", "Coverage % (tape)", "Challenger ECL", "Coverage % (challenger)", "Challenger vs tape"])
rng = lambda col: f"ECL_LoanLevel!${col}$6:${col}${last}"
for i, s in enumerate((1, 2, 3)):
    r = 4 + i
    put(Z, f"A{r}", s, f_in)
    put(Z, f"B{r}", f"=COUNTIF({rng('B')},A{r})", f_n, "#,##0")
    put(Z, f"C{r}", f"=SUMIFS({rng('H')},{rng('B')},A{r})", f_n, "#,##0")
    put(Z, f"D{r}", f"=SUMIFS({rng('I')},{rng('B')},A{r})", f_n, "#,##0")
    put(Z, f"E{r}", f"=SUMIFS({rng('J')},{rng('B')},A{r})", f_n, "#,##0")
    put(Z, f"F{r}", f"=E{r}-D{r}", f_n, "#,##0")
    put(Z, f"G{r}", f"=IF(C{r}=0,0,D{r}/C{r})", f_n, "0.00%")
    put(Z, f"H{r}", f"=SUMIFS({rng('O')},{rng('B')},A{r})", f_n, "#,##0")
    put(Z, f"I{r}", f"=IF(C{r}=0,0,H{r}/C{r})", f_n, "0.00%")
    put(Z, f"J{r}", f"=IF(D{r}=0,0,H{r}/D{r}-1)", f_n, "0.0%")
put(Z, "A7", "Total", f_b)
for col in "BCDEFH":
    put(Z, f"{col}7", f"=SUM({col}4:{col}6)", f_b, "#,##0")
put(Z, "G7", "=IF(C7=0,0,D7/C7)", f_b, "0.00%")
put(Z, "I7", "=IF(C7=0,0,H7/C7)", f_b, "0.00%")
put(Z, "J7", "=IF(D7=0,0,H7/D7-1)", f_b, "0.0%")
put(Z, "A9", "Checks", f_b)
put(Z, "A10", "Tape ECL = PD x LGD x EAD (max loan-level variance < 0.05)")
put(Z, "E10", f"=MAX(MAX({rng('K')}),-MIN({rng('K')}))", f_n, "0.0000")
put(Z, "F10", '=IF(E10<0.05,"PASS","FAIL")', f_b)
put(Z, "A11", "Excel challenger ECL = Python challenger ECL (max loan diff < 0.05)")
put(Z, "E11", f"=MAX(MAX({rng('Q')}),-MIN({rng('Q')}))", f_n, "0.0000")
put(Z, "F11", '=IF(E11<0.05,"PASS","FAIL")', f_b)
put(Z, "A12", "Stage ECL sums to total tape ECL")
put(Z, "E12", f"=SUM({rng('I')})-D7", f_n, "0.00")
put(Z, "F12", '=IF(ABS(E12)<0.5,"PASS","FAIL")', f_b)
put(Z, "A14", "Reading the result", f_b)
notes = [
    "The tape ECL is internally consistent (PD x LGD x EAD), but the challenger is materially higher.",
    "Reason: the tape PD for Stage 1 is about 2%, while the roll-rate matrix in dpd_snapshot_history implies a 12-month PD of about 11% for loans starting 'Current'.",
    "Caveats: only 11 monthly transitions, 31 observed entries into 90+, Markov and time-homogeneity assumptions, and the duplicated-May snapshot (see DataQualityLog DQ05).",
    "Treat the gap as a model-risk finding to discuss in the report, not as a corrected provision.",
]
for i, t in enumerate(notes, start=15):
    put(Z, f"A{i}", t, Font(name=F, size=10))
for col, w in zip("ABCDEFGHIJ", [58, 10, 16, 16, 16, 12, 14, 16, 16, 14]):
    Z.column_dimensions[col].width = w

wb.move_sheet("ECL_Summary", offset=-4)
wb.save("483553A_ManyaSyal_ECL_Waterfall_Validation.xlsx")
print("saved; loans", n, "| order:", wb.sheetnames)
